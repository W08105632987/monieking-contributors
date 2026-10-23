"""Module D1: offline customer import — template, preview, commit, claim by phone, withdrawals, undo, cleanup routes."""
import asyncio
import io
import os

import pytest

TEST_DB_URL = os.environ.get("TEST_DATABASE_URL", "")
_db_name = TEST_DB_URL.rsplit("/", 1)[-1].split("?")[0].lower() if TEST_DB_URL else ""
pytestmark = pytest.mark.skipif(not TEST_DB_URL or "test" not in _db_name, reason="needs TEST_DATABASE_URL (a *test* database)")

if TEST_DB_URL and "test" in _db_name:
    os.environ["DATABASE_URL"] = TEST_DB_URL
    import httpx
    from sqlalchemy import text, select, func
    from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker, AsyncSession
    import app.main as main_mod
    from app.core.database import Base
    from app.core.dependencies import get_current_user
    from app.core.security import hash_withdrawal_password
    from app.models.user import User, UserRole, UserStatus
    from app.models.zone import Zone
    from app.models.card import ContributionCard, ContributionRecord
    from app.models.withdrawal import Withdrawal
    from app.services.offline_import import claim_for_user


def _xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["Full name *", "Phone number *", "Zone *", "Card type *", "Daily rate (₦) *", "Days contributed so far *",
               "Days already withdrawn", "Amount withdrawn (₦)", "Card start date", "Handled by", "Notes"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


async def _run(monkeypatch):
    engine = create_async_engine(TEST_DB_URL)
    async with engine.begin() as conn:
        for stmt in ("DROP SCHEMA public CASCADE", "CREATE SCHEMA public",
                     "CREATE SEQUENCE card_number_seq START 1015", "CREATE SEQUENCE customer_number_seq START 1000"):
            await conn.execute(text(stmt))
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with maker() as db:
        z = Zone(name="Ikeja")
        db.add(z)
        await db.flush()
        d = User(role=UserRole.DIRECTOR, full_name="Dir", phone_number="08040000000", status=UserStatus.ACTIVE,
                 payment_pin_hash=hash_withdrawal_password("4821"))
        t = User(role=UserRole.CUSTOMER, full_name="Test Person", phone_number="08011111111", status=UserStatus.ACTIVE, is_test=True)
        db.add_all([d, t])
        await db.commit()
        dir_id, test_id = d.id, t.id

    from fastapi import Request

    async def override(request: Request):
        async with maker() as db:
            return await db.get(User, dir_id)
    main_mod.app.dependency_overrides[get_current_user] = override

    from app.core.database import get_db

    async def db_override():
        async with maker() as s:
            yield s
    main_mod.app.dependency_overrides[get_db] = db_override

    deleted_logins = []

    async def fake_admin(method, path, **kw):
        deleted_logins.append(path)
        class R:  # noqa: N801
            status_code = 204
        return R()
    monkeypatch.setattr("app.utils.supabase_admin_client.supabase_admin_request", fake_admin)

    transport = httpx.ASGITransport(app=main_mod.app)
    async with httpx.AsyncClient(transport=transport, base_url="http://t") as c:
        base = "/api/v1/zonal-customers"
        r = await c.get(base + "/template.xlsx")
        assert r.status_code == 200 and r.content[:2] == b"PK"

        good = [
            ("Ada Obi", "08031234567", "Ikeja", "Regular", 500, 10, 4, None, "2026-03-02", "Self", None),
            ("Musa Bello", "07068881234", "Ikeja", "Food", 1000, 372, 372, None, None, "Officer", None),
        ]
        bad = good + [("X", "123", "Nowhere", "Gold", -5, 999, None, None, "soon", "Maybe", None)]
        r = await c.post(base + "/import/preview", files={"file": ("a.xlsx", _xlsx(bad))})
        pv = r.json()
        assert r.status_code == 200 and pv["ok"] == 2 and pv["bad"] == 1
        r = await c.post(base + "/import/commit", files={"file": ("a.xlsx", _xlsx(bad))})
        assert r.status_code == 400                                   # nothing saved while a row is wrong
        r = await c.post(base + "/import/commit", files={"file": ("a.xlsx", _xlsx(good))})
        assert r.status_code == 200 and r.json()["staged"] == 2, r.text
        batch_id = r.json()["batch_id"]
        r = await c.post(base + "/import/commit", files={"file": ("a.xlsx", _xlsx(good))})
        assert r.status_code == 400                                   # duplicates are refused

        r = await c.post(base + "/single", json={"full_name": "New Person", "phone_number": "09011112222", "zone_id": (await c.get(base + "/zones")).json()[0]["id"],
                                                 "card_type": "regular", "rate_naira": 200, "days_contributed": 0})
        assert r.status_code == 200, r.text
        assert len((await c.get(base + "/waiting?claimed=false")).json()) == 3

        # Ada signs up with the same phone -> her record is there
        async with maker() as db:
            ada = User(role=UserRole.CUSTOMER, full_name="Ada Obi", phone_number="08031234567", status=UserStatus.ACTIVE)
            db.add(ada)
            await db.flush()
            assert await claim_for_user(db, ada) == 1
            assert await claim_for_user(db, ada) == 0                 # idempotent
            await db.commit()
            card = (await db.execute(select(ContributionCard).where(ContributionCard.owner_id == ada.id))).scalar_one()
            assert card.total_days_contributed == 10 and card.total_contributed_kobo == 6 * 50_000
            recs = (await db.execute(select(ContributionRecord).where(ContributionRecord.card_id == card.id).order_by(ContributionRecord.logical_day))).scalars().all()
            assert len(recs) == 10 and sum(r.is_withdrawn for r in recs) == 4 and recs[0].is_withdrawn and not recs[9].is_withdrawn
            w = (await db.execute(select(Withdrawal).where(Withdrawal.customer_id == ada.id))).scalar_one()
            assert w.status.value == "paid" and w.requested_amount_kobo == 4 * 50_000
            assert ada.zone_id is not None

        r = await c.post(f"{base}/batches/{batch_id}/undo")
        assert r.status_code == 200 and r.json()["removed"] == 1 and r.json()["already_claimed_kept"] == 1   # Musa removed, Ada kept

        # ── cleanup routes ──
        cl = "/api/v1/cleanup"
        r = await c.get(cl + "/users.xlsx")
        assert r.status_code == 200
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Users"]
        marks = {row[1].value: row[0].value for row in ws.iter_rows(min_row=2)}
        assert marks["Dir"] == "Real"
        r = await c.post(cl + "/mark", files={"file": ("u.xlsx", r.content)})
        assert r.status_code == 200, r.text
        s = (await c.get(cl + "/preview")).json()
        assert s["can_purge"] and s["users"] >= 1
        bad = await c.post(cl + "/purge", json={"confirm_phrase": "nope", "pin": "4821", "backup_confirmed": True})
        assert bad.status_code == 400
        bad = await c.post(cl + "/purge", json={"confirm_phrase": "DELETE TEST ACCOUNTS", "pin": "0000", "backup_confirmed": True})
        assert bad.status_code == 403
        r = await c.post(cl + "/purge", json={"confirm_phrase": "DELETE TEST ACCOUNTS", "pin": "4821", "backup_confirmed": True, "acknowledge_value": True})
        assert r.status_code == 200, r.text
        assert any(str(test_id) in p for p in deleted_logins)          # the login was removed too
        async with maker() as db:
            names = {u.full_name for u in (await db.execute(select(User))).scalars().all()}
            assert "Test Person" not in names and "Dir" in names and "Ada Obi" in names
        rc = (await c.get(cl + "/reconcile")).json()
        assert rc["clean"], rc
        r = await c.post(cl + "/reset-stats", json={"confirm_phrase": "RESET STATISTICS", "pin": "4821"})
        assert r.status_code == 200
    main_mod.app.dependency_overrides.clear()
    await engine.dispose()


def test_import_and_cleanup(monkeypatch):
    asyncio.run(_run(monkeypatch))
